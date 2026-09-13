import os
import time
import uuid
import pymupdf  # PyMuPDF
from docx import Document as DocxDocument
from pptx import Presentation
from typing import List, Tuple
from ..schemas.document import DocumentChunk, DocumentMetadata
from ..utils.errors import DocumentProcessingError, UnsupportedFileTypeError
from ..utils.logger import logger
from ..utils.security import sanitize_filename, is_safe_path, validate_uuid
from .metadata_filter import metadata_filter

TEMP_STORAGE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "temp_uploads")
os.makedirs(TEMP_STORAGE_DIR, exist_ok=True)

class DocumentService:
    @staticmethod
    def cleanup_expired_temp_files(max_age_seconds: int = 7200):
        """
        Removes temporary files older than max_age_seconds (default 2 hours)
        to prevent disk exhaustion and data accumulation.
        """
        try:
            now = time.time()
            if not os.path.exists(TEMP_STORAGE_DIR):
                return
            for fname in os.listdir(TEMP_STORAGE_DIR):
                fpath = os.path.join(TEMP_STORAGE_DIR, fname)
                if os.path.isfile(fpath):
                    if now - os.path.getmtime(fpath) > max_age_seconds:
                        try:
                            os.remove(fpath)
                            logger.info(f"Purged expired temporary upload: {fname}")
                        except OSError as e:
                            logger.warning(f"Could not delete expired temp file {fname}: {e}")
        except Exception as e:
            logger.error(f"Error during temp file cleanup: {e}")

    @classmethod
    def save_temp_file(cls, file_bytes: bytes, filename: str) -> Tuple[str, str]:
        clean_name = sanitize_filename(filename)
        ext = os.path.splitext(clean_name)[1].lower()
        if ext not in [".pdf", ".docx", ".pptx", ".txt"]:
            raise UnsupportedFileTypeError()
            
        file_id = str(uuid.uuid4())
        saved_filename = f"{file_id}_{clean_name}"
        saved_path = os.path.join(TEMP_STORAGE_DIR, saved_filename)
        
        # Verify path containment to eliminate path traversal
        if not is_safe_path(TEMP_STORAGE_DIR, saved_path):
            raise DocumentProcessingError(detail="Invalid file path detected.")
            
        with open(saved_path, "wb") as f:
            f.write(file_bytes)
            
        # Run opportunistic cleanup of expired temp files
        cls.cleanup_expired_temp_files()
        
        return file_id, saved_path

    @staticmethod
    def get_file_path_by_id(file_id: str) -> str:
        # Validate file_id structure
        clean_id = (file_id or "").strip()
        if not clean_id or len(clean_id) < 8 or not validate_uuid(clean_id):
            raise DocumentProcessingError(detail="Invalid or malformed document identifier.")
            
        prefix = f"{clean_id}_"
        for f in os.listdir(TEMP_STORAGE_DIR):
            if f.startswith(prefix):
                full_path = os.path.join(TEMP_STORAGE_DIR, f)
                if is_safe_path(TEMP_STORAGE_DIR, full_path) and os.path.isfile(full_path):
                    return full_path
                    
        raise DocumentProcessingError(detail="Uploaded document not found or expired. Please upload again.")

    @classmethod
    def extract_chunks_and_metadata(cls, file_path: str, original_filename: str) -> Tuple[DocumentMetadata, List[DocumentChunk]]:
        if not is_safe_path(TEMP_STORAGE_DIR, file_path) and not os.path.exists(file_path):
            raise DocumentProcessingError(detail="Document file access denied.")

        ext = os.path.splitext(file_path)[1].lower()
        file_size = os.path.getsize(file_path)
        file_id = os.path.basename(file_path).split("_")[0]
        
        chunks: List[DocumentChunk] = []
        total_units = 0
        
        try:
            if ext == ".pdf":
                doc = pymupdf.open(file_path)
                total_units = len(doc)
                logger.info(f"Extracting PDF: {original_filename} ({total_units} pages)")
                
                for page_idx in range(total_units):
                    page = doc[page_idx]
                    text = page.get_text().strip()
                    if text:
                        # Extract first true academic heading line (ignoring metadata, author lines, and slide numbers)
                        lines = [line.strip() for line in text.split("\n") if line.strip()]
                        academic_lines = [l for l in lines if not metadata_filter.is_metadata_line(l)]
                        raw_candidate = academic_lines[0] if academic_lines else (lines[0] if lines else f"Page {page_idx + 1}")
                        page_title = metadata_filter.clean_academic_title(raw_candidate, fallback_content=text)
                        if len(page_title) > 60:
                            page_title = page_title[:60] + "..."
                            
                        chunks.append(DocumentChunk(
                            chunk_id=f"page_{page_idx + 1}",
                            source_reference=f"Page {page_idx + 1}",
                            text=text,
                            page_or_slide_num=page_idx + 1,
                            title=page_title
                        ))
                doc.close()
                
            elif ext == ".docx":
                doc = DocxDocument(file_path)
                current_section = []
                current_heading = "Introduction"
                section_count = 1
                
                for p in doc.paragraphs:
                    t = p.text.strip()
                    if not t:
                        continue
                    if p.style.name.startswith("Heading") or (len(t) < 60 and t.endswith(":")):
                        if current_section:
                            chunks.append(DocumentChunk(
                                chunk_id=f"sec_{section_count}",
                                source_reference=f"Section: {current_heading}",
                                text="\n".join(current_section),
                                page_or_slide_num=section_count,
                                title=current_heading
                            ))
                            section_count += 1
                            current_section = []
                        current_heading = t
                    else:
                        current_section.append(t)
                        
                if current_section:
                    chunks.append(DocumentChunk(
                        chunk_id=f"sec_{section_count}",
                        source_reference=f"Section: {current_heading}",
                        text="\n".join(current_section),
                        page_or_slide_num=section_count,
                        title=current_heading
                    ))
                total_units = max(1, len(chunks))
                logger.info(f"Extracting DOCX: {original_filename} ({total_units} sections)")
                
            elif ext == ".pptx":
                prs = Presentation(file_path)
                total_units = len(prs.slides)
                logger.info(f"Extracting PPTX: {original_filename} ({total_units} slides)")
                
                for slide_idx, slide in enumerate(prs.slides):
                    slide_texts = []
                    slide_title = f"Slide {slide_idx + 1}"
                    
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                pt = paragraph.text.strip()
                                if pt:
                                    slide_texts.append(pt)
                        elif hasattr(shape, "text") and shape.text.strip():
                            slide_texts.append(shape.text.strip())
                            
                    text = "\n".join(slide_texts).strip()
                    if text:
                        candidate_title = None
                        if slide.shapes.title and slide.shapes.title.text.strip():
                            candidate_title = slide.shapes.title.text.strip()
                        elif slide_texts:
                            candidate_title = slide_texts[0][:60]
                        
                        slide_title = metadata_filter.clean_academic_title(candidate_title or f"Slide {slide_idx + 1}", fallback_content=text)
                            
                        chunks.append(DocumentChunk(
                            chunk_id=f"slide_{slide_idx + 1}",
                            source_reference=f"Slide {slide_idx + 1}",
                            text=text,
                            page_or_slide_num=slide_idx + 1,
                            title=slide_title
                        ))
                        
            elif ext == ".txt":
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]
                total_units = len(paragraphs)
                logger.info(f"Extracting TXT: {original_filename} ({total_units} paragraphs)")
                
                for idx, p in enumerate(paragraphs):
                    first_line = p.split("\n")[0][:60]
                    sec_title = metadata_filter.clean_academic_title(first_line or f"Section {idx + 1}", fallback_content=p)
                    chunks.append(DocumentChunk(
                        chunk_id=f"chunk_{idx + 1}",
                        source_reference=f"Section {idx + 1}",
                        text=p,
                        page_or_slide_num=idx + 1,
                        title=sec_title
                    ))
            else:
                raise UnsupportedFileTypeError()
                
            if not chunks:
                raise DocumentProcessingError(detail="The uploaded file appears to be empty or contains unreadable text.")
                
            processed_units = len(chunks)
            total_chars = sum(len(c.text) for c in chunks)
            coverage_pct = 100 if total_units == 0 else int(round((processed_units / total_units) * 100))
            preview = chunks[0].text[:300] + ("..." if len(chunks[0].text) > 300 else "")
            
            logger.info(f"Document Extraction Complete: {original_filename} | Total: {total_units} | Extracted: {processed_units} | Coverage: {coverage_pct}%")
            
            metadata = DocumentMetadata(
                file_id=file_id,
                filename=original_filename,
                file_type=ext.upper().replace(".", ""),
                file_size_bytes=file_size,
                page_count=total_units,
                char_count=total_chars,
                extracted_preview=preview,
                total_units=total_units,
                processed_units=processed_units,
                coverage_percentage=coverage_pct
            )
            return metadata, chunks
            
        except Exception as e:
            logger.error(f"Error processing document {file_path}: {e}", exc_info=True)
            if isinstance(e, (DocumentProcessingError, UnsupportedFileTypeError)):
                raise e
            # Sanitize client-facing error message
            raise DocumentProcessingError(detail="We couldn't process this document. Please check that the file is not corrupted or password-protected.")

document_service = DocumentService()
